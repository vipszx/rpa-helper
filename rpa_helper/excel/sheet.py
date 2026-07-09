import os
from openpyxl import Workbook, load_workbook


class ExcelRow:
    def __init__(self, sheet, row_num, values):
        self.sheet = sheet
        self.row_num = row_num
        self.values = list(values)

    def __getitem__(self, key):
        idx = self.sheet.header_index[key]

        if idx >= len(self.values):
            return None

        return self.values[idx]

    def __setitem__(self, key, value):
        self.sheet.set_value(self.row_num, key, value)

        idx = self.sheet.header_index[key]

        while len(self.values) <= idx:
            self.values.append(None)

        self.values[idx] = value

    def __contains__(self, key):
        return key in self.sheet.header_index

    def get(self, key, default=None):
        try:
            value = self[key]
            return default if value is None else value
        except KeyError:
            return default

    def to_dict(self):
        return {
            header: self.get(header)
            for header in self.sheet.headers
        }

    def __repr__(self):
        return f"<ExcelRow row={self.row_num}>"


class ExcelSheet:
    def __init__(self, file_path):
        self.file_path = file_path

        # 文件不存在 → 创建
        if os.path.exists(file_path):
            self.wb = load_workbook(file_path)
        else:
            self.wb = Workbook()

        self.ws = self.wb.active
        self._refresh_headers()

    def _refresh_headers(self):
        self.headers = [cell.value for cell in self.ws[1]]
        if all(v is None for v in self.headers):
            self.headers = []
            self.header_index = {}
        else:
            self.header_index = {name: idx for idx, name in enumerate(self.headers) if name}

    def _set_active_sheet(self, worksheet):
        self.ws = worksheet
        self.wb.active = self.wb.index(worksheet)
        self._refresh_headers()

    def get_sheet_names(self):
        return self.wb.sheetnames

    def sheet_exists(self, sheet_name):
        return sheet_name in self.wb.sheetnames

    def get_current_sheet_name(self):
        return self.ws.title

    def switch_sheet(self, sheet_name):
        if not self.sheet_exists(sheet_name):
            raise ValueError(f"Sheet not found: {sheet_name}")

        self._set_active_sheet(self.wb[sheet_name])
        return self

    def add_sheet(self, sheet_name, index=None):
        if self.sheet_exists(sheet_name):
            raise ValueError(f"Sheet already exists: {sheet_name}")

        self._set_active_sheet(self.wb.create_sheet(title=sheet_name, index=index))
        return self

    def rename_sheet(self, old_name, new_name):
        if not self.sheet_exists(old_name):
            raise ValueError(f"Sheet not found: {old_name}")
        if old_name != new_name and self.sheet_exists(new_name):
            raise ValueError(f"Sheet already exists: {new_name}")

        self.wb[old_name].title = new_name
        return self

    def delete_sheet(self, sheet_name):
        if not self.sheet_exists(sheet_name):
            raise ValueError(f"Sheet not found: {sheet_name}")
        if len(self.wb.sheetnames) == 1:
            raise ValueError("Cannot delete the only sheet")

        delete_index = self.wb.sheetnames.index(sheet_name)
        current_sheet_name = self.ws.title
        self.wb.remove(self.wb[sheet_name])

        if sheet_name == current_sheet_name:
            next_index = min(delete_index, len(self.wb.worksheets) - 1)
            self._set_active_sheet(self.wb.worksheets[next_index])
        else:
            self._set_active_sheet(self.wb[current_sheet_name])
        return self

    def has_column(self, name):
        return name in self.header_index

    def get_columns(self):
        return self.headers

    def add_column(self, name):
        if self.has_column(name):
            return self.header_index[name]

        col_idx = len(self.headers)
        self.ws.cell(row=1, column=col_idx + 1, value=name)
        self.headers.append(name)
        self.header_index[name] = col_idx

        return col_idx

    def ensure_columns(self, *columns):
        for col in columns:
            self.add_column(col)

    def require_columns(self, *columns):
        missing = [
            col
            for col in columns
            if col not in self.header_index
        ]

        return missing

    def clear_sheet(self, keep_header=True):
        if keep_header:
            if self.ws.max_row > 1:
                self.ws.delete_rows(2, self.ws.max_row - 1)
        else:
            self.ws.delete_rows(1, self.ws.max_row)

        self._refresh_headers()
        return self

    def get_value(self, row, field, default=None):
        idx = self.header_index.get(field)
        if idx is None or idx >= len(row):
            return default
        return row[idx]

    def set_value(self, row_num, field, value):
        if not self.has_column(field):
            self.add_column(field)
        col_idx = self.header_index[field] + 1
        self.ws.cell(row=row_num, column=col_idx, value=value)

    def append(self, data: dict):
        for key in data.keys():
            if key not in self.header_index:
                self.add_column(key)

        row = []
        for header in self.headers:
            row.append(data.get(header))
        self.ws.append(row)

    def iter_rows(self):
        for row_num, values in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2):
            yield row_num, values

    def rows(self):
        for row_num, values in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2):
            yield ExcelRow(self, row_num, values)

    def save(self, file_path=None):
        self.wb.save(file_path or self.file_path)

    def close(self):
        self.wb.close()

if __name__ == "__main__":
    excel = ExcelSheet(r"C:\Users\Administrator\Desktop\test.xlsx")
    for row in excel.rows():
        print(row.to_dict())
    # excel.append({"name": "test", "value": 1})
    # excel.save()
